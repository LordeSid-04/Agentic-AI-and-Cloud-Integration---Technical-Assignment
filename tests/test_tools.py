import unittest
from pathlib import Path
import openpyxl
from decimal import Decimal, ROUND_HALF_UP

from data_loader import load_energy_data
from tools import _parse_range, get_total_consumption, generate_invoice
from config import TARIFF_SGD_PER_KWH


class TestEnergyTools(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.df, cls.min_date, cls.max_date = load_energy_data()

    def test_01_load_energy_data(self):
        """Test dataset loading, row count, and coverage bounds."""
        self.assertEqual(len(self.df), 744)
        self.assertEqual(str(self.min_date.date()), "2026-08-01")
        self.assertEqual(str(self.max_date.date()), "2026-08-31")
        self.assertTrue((self.df["energy_consumption_kwh"] >= 0).all())

    def test_02_parse_range_valid(self):
        """Test valid YYYY-MM-DD date range parsing."""
        start, end, err = _parse_range("2026-08-08", "2026-08-14")
        self.assertIsNone(err)
        self.assertEqual(str(start.date()), "2026-08-08")
        self.assertEqual(str(end.date()), "2026-08-14")

    def test_03_parse_range_invalid(self):
        """Test error handling for missing, unparseable, and reversed dates."""
        _, _, err_missing = _parse_range("", "2026-08-14")
        self.assertIn("required", err_missing)

        _, _, err_unparseable = _parse_range("invalid-date", "2026-08-14")
        self.assertIn("Could not parse", err_unparseable)

        _, _, err_reversed = _parse_range("2026-08-14", "2026-08-08")
        self.assertIn("after end date", err_reversed)

    def test_04_get_total_consumption_normal(self):
        """Test consumption calculation for a standard 7-day period (168 hours)."""
        result = get_total_consumption(self.df, "2026-08-08", "2026-08-14")
        self.assertNotIn("error", result)
        self.assertEqual(result["hours_found"], 168)
        self.assertEqual(result["hours_expected"], 168)
        self.assertEqual(result["total_kwh"], 424.27)
        self.assertNotIn("warning", result)

    def test_05_get_total_consumption_full_month(self):
        """Test consumption calculation for the entire month (744 hours)."""
        result = get_total_consumption(self.df, "2026-08-01", "2026-08-31")
        self.assertNotIn("error", result)
        self.assertEqual(result["hours_found"], 744)
        self.assertEqual(result["hours_expected"], 744)
        self.assertEqual(result["total_kwh"], 1836.50)

    def test_06_get_total_consumption_out_of_range(self):
        """Test that fully out-of-range queries are refused with boundary info."""
        result = get_total_consumption(self.df, "2026-09-01", "2026-09-05")
        self.assertIn("error", result)
        self.assertIn("No data available", result["error"])
        self.assertIn("covers", result["error"])

    def test_07_get_total_consumption_partial_range(self):
        """Test partially covered range warning and hours count."""
        result = get_total_consumption(self.df, "2026-08-28", "2026-09-03")
        self.assertNotIn("error", result)
        self.assertEqual(result["hours_found"], 96)  # 4 days in August * 24 hours
        self.assertEqual(result["covered_start"], "2026-08-28")
        self.assertEqual(result["covered_end"], "2026-08-31")
        self.assertIn("warning", result)
        self.assertIn("Only 96 of 168", result["warning"])

    def test_08_money_rounding_half_up(self):
        """Test that exact half-cents round up (459.125 -> 459.13, not 459.12)."""
        total_kwh = 1836.50
        tariff = TARIFF_SGD_PER_KWH  # 0.25
        # Standard Python round() would fail here with 459.12:
        py_bankers_round = round(total_kwh * tariff, 2)
        self.assertEqual(py_bankers_round, 459.12)

        # Decimal half-up produces the correct financial invoice amount:
        decimal_half_up = float(
            (Decimal(str(total_kwh)) * Decimal(str(tariff)))
            .quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        )
        self.assertEqual(decimal_half_up, 459.13)

        # Verify generate_invoice uses half-up rounding:
        invoice = generate_invoice(self.df, "2026-08-01", "2026-08-31")
        self.assertEqual(invoice["total_payable_sgd"], 459.13)

    def test_09_generate_invoice_output_structure(self):
        """Test invoice structure: summary fields, daily table columns, and exact sums."""
        invoice = generate_invoice(self.df, "2026-08-01", "2026-08-31")
        self.assertTrue(Path(invoice["file_path"]).exists())
        self.assertEqual(invoice["billing_period"], "2026-08-01 to 2026-08-31")
        self.assertEqual(invoice["total_kwh"], 1836.50)
        self.assertEqual(invoice["total_payable_sgd"], 459.13)

        # Verify Excel file contents
        wb = openpyxl.load_workbook(invoice["file_path"])
        ws = wb["Invoice"]
        # Breakdown header has only Date and Consumption (kWh), no Cost column
        header_row = [c.value for c in ws[7]]
        self.assertEqual(header_row, ["Date", "Consumption (kWh)"])

        # Verify daily kWh sums exactly to total kWh
        daily_kwh_values = [ws.cell(row=r, column=2).value for r in range(8, 39)]
        self.assertEqual(len(daily_kwh_values), 31)
        self.assertEqual(round(sum(daily_kwh_values), 2), 1836.50)

    def test_10_generate_invoice_cell_number_formatting(self):
        """Test that numeric cells have '0.00' number format in openpyxl."""
        invoice = generate_invoice(self.df, "2026-08-01", "2026-08-31")
        wb = openpyxl.load_workbook(invoice["file_path"])
        ws = wb["Invoice"]

        # Summary values formatting
        self.assertEqual(ws["B3"].number_format, "0.00")  # Total Consumption
        self.assertEqual(ws["B4"].number_format, "0.00")  # Tariff
        self.assertEqual(ws["B5"].number_format, "0.00")  # Total Payable

        # Daily breakdown values formatting
        self.assertEqual(ws.cell(row=8, column=2).number_format, "0.00")
        self.assertEqual(ws.cell(row=38, column=2).number_format, "0.00")


if __name__ == "__main__":
    unittest.main()
